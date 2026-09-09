"""Outbound campaign drafts. Telephony and transcript extraction are intentionally gated."""
import csv
import io
import json
import re
import sqlite3
import uuid
import zipfile
from datetime import datetime, timezone
from contextlib import contextmanager
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel, Field
import auth
from outbound_settings import CallingSettings, CallOptions, protect_token

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
PERMISSIONS = ('outbound.view', 'outbound.manage', 'outbound.export', 'outbound.configure', 'outbound.voice')


def require_permission(key):
    def check(credentials=Depends(auth.security)):
        user = auth._authenticate(credentials)
        if not user['is_superadmin'] and not auth._has_permission(user['id'], key):
            raise HTTPException(403, 'Permission required: ' + key)
        return user['username']
    return check


def require_voice_access(app_key, credentials, available_bots):
    user = auth._authenticate(credentials)
    if not user['is_superadmin'] and not auth._has_permission(user['id'], 'outbound.voice'):
        raise HTTPException(403, 'Permission required: outbound.voice')
    if app_key not in available_bots:
        raise HTTPException(404, 'Unknown voice bot')
    if not user['is_superadmin'] and not auth._has_permission(user['id'], app_key):
        raise HTTPException(403, 'Access to the selected customer care bot is also required')
    return user['username']


class Question(BaseModel):
    field: str = Field(min_length=1, max_length=80, pattern=r'^[a-zA-Z][a-zA-Z0-9_]*$')
    question: str = Field(min_length=1, max_length=500)
    type: str = Field(default='text', pattern=r'^(text|yes_no|number|datetime)$')


class Campaign(BaseModel):
    calling: CallOptions = Field(default_factory=CallOptions)
    name: str = Field(min_length=1, max_length=150)
    agent: str = Field(min_length=1, max_length=80)
    purpose: str = Field(default='', max_length=2000)
    opening: str = Field(default='', max_length=2000)
    instructions: str = Field(default='', max_length=6000)
    sheet_url: str = Field(default='', max_length=500)
    questions: list[Question] = Field(default_factory=list, max_length=30)


class LeadImport(BaseModel):
    headers: list[str] = Field(max_length=100)
    rows: list[list[str]] = Field(max_length=MAX_ROWS)
    phone: str
    name: str = ''
    purpose: str = ''
    country_code: str = Field(default='', max_length=4, pattern=r'^\+?\d{0,3}$')


def parse_upload(data, filename):
    if len(data) > MAX_BYTES:
        raise HTTPException(413, 'Maximum upload size is 5 MB')
    try:
        if filename.lower().endswith('.csv'):
            table = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        elif filename.lower().endswith('.xlsx'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('Expanded Excel file exceeds 30 MB')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                table = []
                for row in workbook.worksheets[0].iter_rows(values_only=True):
                    if len(table) > MAX_ROWS or len(row) > 100:
                        raise ValueError('Maximum 5,000 rows and 100 columns')
                    table.append(['' if v is None else str(v) for v in row])
            finally:
                workbook.close()
        else:
            raise ValueError('Use a UTF-8 CSV or XLSX file')
        table = [row for row in table if any(str(v).strip() for v in row)]
        if not table or len(table) < 2:
            raise ValueError('Include a header and at least one lead')
        headers = [str(v).strip() for v in table[0]]
        if len(headers) > 100 or len(table) - 1 > MAX_ROWS:
            raise ValueError('Maximum 5,000 rows and 100 columns')
        if any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('Column headings must be non-empty and unique')
        if any(len(row) > len(headers) for row in table[1:]):
            raise ValueError('Some rows have more cells than the header')
        rows = [row + [''] * (len(headers) - len(row)) for row in table[1:]]
        aliases = {'phone': ['phone', 'phonenumber', 'mobile', 'mobilenumber', 'contactnumber'],
                   'name': ['name', 'customername', 'fullname'], 'purpose': ['purpose', 'callpurpose', 'reason']}
        suggested = {}
        for key, names in aliases.items():
            candidates = [h for h in headers if re.sub(r'[^a-z]', '', h.lower()) in names]
            suggested[key] = candidates[0] if len(candidates) == 1 else ''
        return {'headers': headers, 'rows': rows, 'suggested': suggested}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, str(exc) if isinstance(exc, ValueError) else 'Unable to read this file') from exc


def normalize_phone(raw, code):
    value = re.sub(r'[\s().-]', '', raw.strip())
    if value.startswith('00'):
        value = '+' + value[2:]
    if not value.startswith('+') and code:
        value = '+' + code.lstrip('+') + value
    return value if re.fullmatch(r'\+[1-9]\d{7,14}', value) else ''


def csv_cell(value):
    text = str(value or '')
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')) else text


def create_router(db_path: Path, static_dir: Path, list_agents):
    router = APIRouter()
    view = require_permission('outbound.view')
    manage = require_permission('outbound.manage')
    export = require_permission('outbound.export')
    configure = require_permission('outbound.configure')

    @contextmanager
    def connect():
        con = sqlite3.connect(db_path)
        con.row_factory = sqlite3.Row
        con.executescript('''
            CREATE TABLE IF NOT EXISTS calling_settings(id INTEGER PRIMARY KEY CHECK(id=1), config TEXT NOT NULL, protected_token TEXT NOT NULL DEFAULT '');
            CREATE TABLE IF NOT EXISTS campaigns(id TEXT PRIMARY KEY, config TEXT NOT NULL, updated TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS leads(id TEXT PRIMARY KEY, campaign_id TEXT NOT NULL, phone TEXT NOT NULL,
                name TEXT NOT NULL, purpose TEXT NOT NULL, original TEXT NOT NULL,
                UNIQUE(campaign_id, phone));
            CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, actor TEXT, action TEXT, campaign_id TEXT, at TEXT);
        ''')
        try:
            with con:
                yield con
        finally:
            con.close()

    def record(con, actor, action, cid):
        con.execute('INSERT INTO audit(actor, action, campaign_id, at) VALUES(?,?,?,?)',
                    (actor, action, cid, datetime.now(timezone.utc).isoformat()))

    def get_campaign(con, cid):
        row = con.execute('SELECT * FROM campaigns WHERE id=?', (cid,)).fetchone()
        if not row:
            raise HTTPException(404, 'Campaign not found')
        return dict(row)

    def read_settings(con):
        row = con.execute('SELECT config, protected_token FROM calling_settings WHERE id=1').fetchone()
        if row:
            return json.loads(row['config']), row['protected_token']
        return CallingSettings().model_dump(exclude={'auth_token', 'clear_auth_token'}), ''

    @router.get('/admin/api/outbound/calling/catalog')
    def calling_catalog(username=Depends(view)):
        with connect() as con:
            config, _ = read_settings(con)
            return {'provider': config['provider'], 'numbers': config['numbers'],
                    'default_number_id': config['default_number_id'], 'connected': False}

    @router.get('/admin/api/outbound/calling/settings')
    def calling_settings(username=Depends(configure)):
        with connect() as con:
            config, token = read_settings(con)
            return {**config, 'auth_token_set': bool(token), 'connected': False}

    @router.put('/admin/api/outbound/calling/settings')
    def save_calling_settings(data: CallingSettings, username=Depends(configure)):
        with connect() as con:
            con.execute('BEGIN IMMEDIATE')
            previous, token = read_settings(con)
            if data.revision != previous['revision']:
                raise HTTPException(409, 'Calling settings changed in another session. Reload settings before saving.')
            incoming = {n.id: n for n in data.numbers}
            old_numbers = {n['id']: n for n in previous['numbers']}
            for row in con.execute('SELECT config FROM campaigns'):
                selected = json.loads(row['config']).get('calling', {}).get('caller_number_id', '')
                if selected and (selected not in incoming or (selected in old_numbers and incoming[selected].phone != old_numbers[selected]['phone'])):
                    raise HTTPException(409, 'A campaign uses this caller number. Reassign that campaign before removing or changing the number.')
            new_token = data.auth_token.get_secret_value() if data.auth_token else ''
            if data.account_sid != previous['account_sid'] and token and not new_token and not data.clear_auth_token:
                raise HTTPException(400, 'Replace or remove the saved auth token when changing the account SID.')
            if data.clear_auth_token:
                token = ''
            elif new_token:
                try:
                    token = protect_token(new_token)
                except RuntimeError as exc:
                    raise HTTPException(503, str(exc)) from exc
            config = data.model_dump(exclude={'auth_token', 'clear_auth_token'})
            config['revision'] += 1
            con.execute('INSERT INTO calling_settings VALUES(1,?,?) ON CONFLICT(id) DO UPDATE SET config=excluded.config,protected_token=excluded.protected_token',
                        (json.dumps(config), token))
            record(con, username, 'save_calling_settings', '')
            return {**config, 'auth_token_set': bool(token), 'connected': False}

    @router.get('/outbound')
    def page(username=Depends(view)):
        return FileResponse(static_dir / 'outbound.html')

    @router.get('/outbound/assets/{name}')
    def asset(name: str, username=Depends(view)):
        if name not in ('outbound.css', 'outbound.js', 'outbound-config.js'):
            raise HTTPException(404)
        return FileResponse(static_dir / name)

    @router.get('/admin/api/outbound/session')
    def session(credentials=Depends(auth.security)):
        user = auth._authenticate(credentials)
        return {'username': user['username'], 'permissions': [p for p in PERMISSIONS
                if user['is_superadmin'] or auth._has_permission(user['id'], p)]}

    @router.get('/admin/api/outbound/agents')
    def agents(username=Depends(view)):
        return {'agents': [{'slug': b['slug'], 'label': b['label']} for b in list_agents()]}

    @router.get('/admin/api/outbound/campaigns')
    def campaigns(username=Depends(view)):
        with connect() as con:
            return {'campaigns': [{'id': r['id'], **json.loads(r['config']), 'updated': r['updated'],
                    'lead_count': r['lead_count'], 'status': 'Draft'} for r in con.execute(
                    'SELECT c.*, (SELECT COUNT(*) FROM leads l WHERE l.campaign_id=c.id) lead_count FROM campaigns c ORDER BY updated DESC')]}

    @router.put('/admin/api/outbound/campaigns/{cid}')
    def save(cid: str, data: Campaign, username=Depends(manage)):
        if not re.fullmatch(r'[a-zA-Z0-9-]{1,80}', cid):
            raise HTTPException(400, 'Invalid campaign ID')
        if data.agent not in {b['slug'] for b in list_agents()}:
            raise HTTPException(400, 'Select an existing customer care bot')
        fields = [q.field.lower() for q in data.questions]
        if len(fields) != len(set(fields)):
            raise HTTPException(400, 'Answer field names must be unique')
        if not data.name.strip():
            raise HTTPException(400, 'Enter a campaign name')
        with connect() as con:
            con.execute('BEGIN IMMEDIATE')
            settings, _ = read_settings(con)
            selected = data.calling.caller_number_id
            existing = con.execute('SELECT config FROM campaigns WHERE id=?', (cid,)).fetchone()
            old_calling = json.loads(existing['config']).get('calling', {}) if existing else {}
            if 'calling' not in data.model_fields_set and existing:
                data.calling = CallOptions(**old_calling)
                selected = data.calling.caller_number_id
            if selected:
                number = next((n for n in settings['numbers'] if n['id'] == selected), None)
                if not number or (not number['enabled'] and old_calling.get('caller_number_id') != selected):
                    raise HTTPException(400, 'Select an enabled caller number from Calling settings.')
            con.execute('INSERT INTO campaigns VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET config=excluded.config,updated=excluded.updated',
                        (cid, data.model_dump_json(), datetime.now(timezone.utc).isoformat()))
            record(con, username, 'save_draft', cid)
        return {'id': cid, 'status': 'Draft'}

    @router.post('/admin/api/outbound/preview')
    async def preview(file: UploadFile = File(...), username=Depends(manage)):
        return parse_upload(await file.read(MAX_BYTES + 1), file.filename or '')

    @router.get('/admin/api/outbound/campaigns/{cid}/leads')
    def leads(cid: str, username=Depends(view)):
        with connect() as con:
            get_campaign(con, cid)
            return {'leads': [dict(r) for r in con.execute('SELECT id,phone,name,purpose FROM leads WHERE campaign_id=?', (cid,))]}

    @router.post('/admin/api/outbound/campaigns/{cid}/leads')
    def import_leads(cid: str, data: LeadImport, username=Depends(manage)):
        if len(set(data.headers)) != len(data.headers) or data.phone not in data.headers:
            raise HTTPException(400, 'Select a valid, unique phone column')
        if any(v and v not in data.headers for v in (data.name, data.purpose)):
            raise HTTPException(400, 'Invalid column mapping')
        if any(len(row) != len(data.headers) or any(len(cell) > 10000 for cell in row) for row in data.rows):
            raise HTTPException(400, 'Invalid row dimensions or oversized cells')
        result = {'imported': 0, 'duplicates': 0, 'rejected': []}
        with connect() as con:
            config = json.loads(get_campaign(con, cid)['config'])
            for index, row in enumerate(data.rows, 2):
                source = dict(zip(data.headers, row))
                phone = normalize_phone(source[data.phone], data.country_code)
                purpose = source.get(data.purpose, '').strip() or config['purpose'].strip()
                if not phone or not purpose:
                    result['rejected'].append({'row': index, 'reason': 'Invalid/incomplete international phone number' if not phone else 'Missing call purpose'})
                    continue
                cur = con.execute('INSERT OR IGNORE INTO leads VALUES(?,?,?,?,?,?)',
                    (str(uuid.uuid4()), cid, phone, source.get(data.name, ''), purpose, json.dumps(source)))
                result['imported' if cur.rowcount else 'duplicates'] += 1
            record(con, username, 'import_leads', cid)
        return result

    @router.get('/admin/api/outbound/campaigns/{cid}/export')
    def download(cid: str, username=Depends(export)):
        with connect() as con:
            config = json.loads(get_campaign(con, cid)['config'])
            rows = list(con.execute('SELECT * FROM leads WHERE campaign_id=?', (cid,)))
            original_headers = list(dict.fromkeys(k for r in rows for k in json.loads(r['original'])))
            output = io.StringIO(newline='')
            writer = csv.writer(output)
            writer.writerow(['Lead ID', 'Customer name', 'Phone', 'Purpose', 'Call status'] +
                            ['Source: ' + csv_cell(k) for k in original_headers] + ['Answer: ' + q['field'] for q in config['questions']])
            for row in rows:
                original = json.loads(row['original'])
                writer.writerow([csv_cell(row[k]) for k in ('id', 'name', 'phone', 'purpose')] + ['Not called'] +
                                [csv_cell(original.get(k, '')) for k in original_headers] + [''] * len(config['questions']))
            record(con, username, 'export_leads', cid)
        return Response('\ufeff' + output.getvalue(), media_type='text/csv',
                        headers={'Content-Disposition': 'attachment; filename="outbound-results.csv"'})

    return router
