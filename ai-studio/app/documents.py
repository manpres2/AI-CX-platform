"""Document ingestion and generation for AI Studio's GPT Chat.

Two directions:
  * IN  — a PDF/Word/text file is turned into plain text the model can read,
          and an image is base64'd for a vision-capable model to look at.
  * OUT — a model reply (loosely markdown) is rendered to a real .pdf/.docx/.txt
          the admin can download.

PDFs are generated with PyMuPDF's Story API and Word files with python-docx —
both already installed for the voice bots' knowledge-base ingestion, so this
adds no new dependency.
"""

import base64
import html
import io
import re
from pathlib import Path

import docx as python_docx
import fitz  # PyMuPDF

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
DOC_EXTENSIONS = {".pdf", ".docx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | DOC_EXTENSIONS | IMAGE_EXTENSIONS

MAX_UPLOAD_BYTES = 20 * 1024 * 1024


def _flatten_table(rows: list[list[str]]) -> str:
    rows = [[(c or "").strip().replace("\n", " ") for c in r] for r in rows if r]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    if not body:
        return " | ".join(c for c in header if c)
    lines = []
    for row in body:
        pairs = [f"{h}: {v}" for h, v in zip(header, row) if h or v]
        if pairs:
            lines.append(" | ".join(pairs))
    return "\n".join(lines)


def _extract_pdf(data: bytes) -> str:
    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            table_rects = []
            try:
                for tbl in page.find_tables().tables:
                    flat = _flatten_table(tbl.extract())
                    if flat:
                        parts.append(flat)
                    table_rects.append(fitz.Rect(tbl.bbox))
            except Exception:
                pass
            words = page.get_text("words")
            kept = [w for w in words if not any(fitz.Rect(w[:4]).intersects(r) for r in table_rects)]
            kept.sort(key=lambda w: (w[5], w[6], w[7]))
            page_text = " ".join(w[4] for w in kept).strip()
            if page_text:
                parts.append(page_text)
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    parts = []
    d = python_docx.Document(io.BytesIO(data))
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in d.tables:
        flat = _flatten_table([[cell.text for cell in row.cells] for row in table.rows])
        if flat:
            parts.append(flat)
    return "\n\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    raise ValueError(f"Cannot extract text from '{ext}' files")


def to_data_uri(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    mime = "jpeg" if ext == "jpg" else ext
    return f"data:image/{mime};base64,{base64.b64encode(data).decode()}"


# ── Generation (model reply → downloadable document) ─────────────────────────
_INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
_INLINE_ITALIC = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)")
_INLINE_CODE = re.compile(r"`([^`]+?)`")


def _inline_html(text: str) -> str:
    """Escape first, then re-introduce only the inline markup we support — so
    model output containing raw HTML can never inject markup into the PDF."""
    out = html.escape(text)
    out = _INLINE_BOLD.sub(r"<b>\1</b>", out)
    out = _INLINE_ITALIC.sub(r"<i>\1</i>", out)
    out = _INLINE_CODE.sub(r"<code>\1</code>", out)
    return out


def _parse_blocks(markdown: str) -> list[tuple[str, str]]:
    """Flattens loose markdown into (kind, text) blocks, where kind is one of
    h1/h2/h3, bullet, number, or para."""
    blocks = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            blocks.append((f"h{len(heading.group(1))}", heading.group(2).strip()))
            continue
        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if bullet:
            blocks.append(("bullet", bullet.group(1).strip()))
            continue
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if numbered:
            blocks.append(("number", numbered.group(1).strip()))
            continue
        blocks.append(("para", line.strip()))
    return blocks


def _blocks_to_html(blocks: list[tuple[str, str]], title: str) -> str:
    body = []
    list_open = None
    for kind, text in blocks:
        wanted = "ul" if kind == "bullet" else "ol" if kind == "number" else None
        if list_open and list_open != wanted:
            body.append(f"</{list_open}>")
            list_open = None
        if wanted and not list_open:
            body.append(f"<{wanted}>")
            list_open = wanted
        if wanted:
            body.append(f"<li>{_inline_html(text)}</li>")
        elif kind.startswith("h"):
            body.append(f"<{kind}>{_inline_html(text)}</{kind}>")
        else:
            body.append(f"<p>{_inline_html(text)}</p>")
    if list_open:
        body.append(f"</{list_open}>")
    return f"""<html><head><style>
      body {{ font-family: sans-serif; font-size: 11pt; line-height: 1.5; }}
      h1 {{ font-size: 19pt; margin: 0 0 10pt; }}
      h2 {{ font-size: 15pt; margin: 14pt 0 6pt; }}
      h3 {{ font-size: 12.5pt; margin: 12pt 0 4pt; }}
      p {{ margin: 0 0 8pt; }}
      li {{ margin: 0 0 4pt; }}
      code {{ font-family: monospace; }}
    </style></head><body><h1>{html.escape(title)}</h1>{''.join(body)}</body></html>"""


def make_pdf(markdown: str, title: str) -> bytes:
    story = fitz.Story(html=_blocks_to_html(_parse_blocks(markdown), title))
    writer = fitz.DocumentWriter(buf := io.BytesIO())
    page_rect = fitz.paper_rect("a4")
    content_rect = page_rect + (54, 54, -54, -54)
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buf.getvalue()


_INLINE_SPLIT = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")


def _add_runs(paragraph, text: str):
    """Splits one line into python-docx runs so **bold**/*italic*/`code` render
    as real Word formatting instead of showing their literal markers."""
    for piece in _INLINE_SPLIT.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def make_docx(markdown: str, title: str) -> bytes:
    doc = python_docx.Document()
    doc.add_heading(title, level=0)
    for kind, text in _parse_blocks(markdown):
        if kind.startswith("h"):
            doc.add_heading(text, level=int(kind[1]))
            continue
        style = {"bullet": "List Bullet", "number": "List Number"}.get(kind)
        _add_runs(doc.add_paragraph(style=style) if style else doc.add_paragraph(), text)
    doc.save(buf := io.BytesIO())
    return buf.getvalue()


def make_document(markdown: str, title: str, fmt: str) -> tuple[bytes, str]:
    """Returns (file bytes, mime type) for fmt in pdf/docx/txt."""
    if fmt == "pdf":
        return make_pdf(markdown, title), "application/pdf"
    if fmt == "docx":
        return make_docx(markdown, title), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if fmt == "txt":
        return f"{title}\n\n{markdown}".encode("utf-8"), "text/plain; charset=utf-8"
    raise ValueError(f"Unsupported format '{fmt}'")
